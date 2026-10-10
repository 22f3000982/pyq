import io
import openpyxl
from backend.models import db,User,Course,Term,ExamType,Paper,Question,QuestionOption,QuestionImage,PaperProgress,Attempt,AttemptAnswer,Bookmark,IngestionBatch,IngestionFile,SourceEntry
from conftest import login


def workbook():
    wb=openpyxl.Workbook()
    master=wb.active;master.title='Courses'
    master.append(['Course Name','Course Code','Course Level','Course Type'])
    master.append(['Deep Learning','DL','Diploma','Theory'])
    q1=wb.create_sheet('May 2026');q1.append(['Course Name','Quiz 1','Quiz 2','FN End Term','AN End Term'])
    for row,name in enumerate(['Deep Learning'],2):
        q1.cell(row,1,name)
        for col,(label,key) in enumerate([('Q1','q1'),('Q2','q2'),('FN','fn'),('AN','an')],2):
            q1.cell(row,col,label);q1.cell(row,col).hyperlink=f'https://drive.google.com/file/d/{key}/view'
    jan=wb.create_sheet('Jan 2026');jan.append(['Course Name','Quiz 1'])
    jan.append(['Deep Learning','Old Q1']);jan['B2'].hyperlink='https://drive.google.com/file/d/q1-old/view'
    out=io.BytesIO();wb.save(out);wb.close();return out.getvalue()


def post_xlsx(client,url,h,data):
    return client.post(url,headers=h,data={'file':(io.BytesIO(data),'catalog.xlsx')},content_type='multipart/form-data')


def test_fresh_catalog_refresh_and_campaign_order(app,client):
    h=login(client,True)
    r=post_xlsx(client,'/api/admin/catalog/refresh',h,workbook())
    assert r.status_code==201
    assert r.json['queued']==0
    campaign=client.get('/api/admin/catalog/campaign',headers=h).json
    assert campaign['total']==5
    assert campaign['current']['label']=='Quiz 1'
    assert campaign['current']['term']=='May 2026'
    labels=[(g['label'],g['term']) for g in campaign['groups']]
    assert labels[:2]==[('Quiz 1','May 2026'),('Quiz 1','Jan 2026')]
    assert ('Quiz 2','May 2026') in labels
    assert ('End Term · FN','May 2026') in labels
    assert ('End Term · AN','May 2026') in labels


def test_campaign_processes_only_current_group(app,client):
    h=login(client,True);post_xlsx(client,'/api/admin/catalog/refresh',h,workbook())
    campaign=client.get('/api/admin/catalog/campaign',headers=h).json
    current=campaign['current']
    r=client.post('/api/admin/catalog/campaign/process',headers=h,json={'stage':current['stage'],'term_id':current['term_id'],'limit':20})
    assert r.status_code==202 and r.json['queued']==1
    queued=IngestionFile.query.all()
    assert len(queued)==1
    p=db.session.get(Paper,queued[0].paper_id)
    assert p.exam_type.name=='Quiz 1' and p.term.name=='May 2026'


def test_controlled_reset_keeps_users_and_clears_library(app,client):
    h=login(client,True);post_xlsx(client,'/api/admin/catalog/refresh',h,workbook())
    p=Paper.query.first();q=Question(paper_id=p.id,number='1',kind='MCQ',text='Q',answers=['A'],answer_status='ANSWER_AVAILABLE',status='AVAILABLE')
    db.session.add(q);db.session.flush()
    db.session.add(QuestionOption(question_id=q.id,key='A',text='A',position=0))
    db.session.add(QuestionImage(question_id=q.id,path='asset.png'))
    student=User(email='s@example.test',name='S',password_hash='x');db.session.add(student);db.session.flush()
    db.session.add(PaperProgress(user_id=student.id,paper_id=p.id,attempted=True,last_score=50,last_attempted_at=1))
    a=Attempt(user_id=student.id,paper_id=p.id,mode='exam',title='x',expires_at=9999999999);db.session.add(a);db.session.flush()
    db.session.add(AttemptAnswer(attempt_id=a.id,question_id=q.id,position=1,snapshot={'id':q.id}))
    db.session.add(Bookmark(user_id=student.id,question_id=q.id));db.session.commit()
    preview=client.get('/api/admin/library-reset/preview',headers=h)
    assert preview.status_code==200 and preview.json['counts']['papers']==5
    bad=client.post('/api/admin/library-reset',headers=h,json={'confirmation':'wrong','cleanup_storage':False})
    assert bad.status_code==400
    done=client.post('/api/admin/library-reset',headers=h,json={'confirmation':'RESET PYQ LIBRARY','cleanup_storage':False})
    assert done.status_code==200
    assert User.query.count()==2
    assert Paper.query.count()==Question.query.count()==Course.query.count()==0
    assert Attempt.query.count()==PaperProgress.query.count()==Bookmark.query.count()==0


def test_reset_preview_marks_old_processing_as_stale_and_allows_cancel(app,client):
    h=login(client,True);post_xlsx(client,'/api/admin/catalog/refresh',h,workbook())
    p=Paper.query.first()
    b=IngestionBatch(user_id=User.query.filter_by(role='ADMIN').first().id);db.session.add(b);db.session.flush()
    f=IngestionFile(batch_id=b.id,paper_id=p.id,filename='stuck.pdf',status='PROCESSING',started_at=1,source_url=p.source_url)
    db.session.add(f);db.session.commit()
    preview=client.get('/api/admin/library-reset/preview',headers=h)
    assert preview.status_code==200
    assert preview.json['counts']['active_ingestion']==1
    assert preview.json['counts']['stale_ingestion']==1
    assert preview.json['counts']['live_ingestion']==0
    done=client.post('/api/admin/library-reset',headers=h,json={'confirmation':'RESET PYQ LIBRARY','cleanup_storage':False,'cancel_active':True})
    assert done.status_code==200
    assert Paper.query.count()==0 and IngestionFile.query.count()==0


def test_live_processing_requires_explicit_cancel(app,client):
    import time
    h=login(client,True);post_xlsx(client,'/api/admin/catalog/refresh',h,workbook())
    p=Paper.query.first()
    b=IngestionBatch(user_id=User.query.filter_by(role='ADMIN').first().id);db.session.add(b);db.session.flush()
    db.session.add(IngestionFile(batch_id=b.id,paper_id=p.id,filename='live.pdf',status='PROCESSING',started_at=time.time(),source_url=p.source_url));db.session.commit()
    blocked=client.post('/api/admin/library-reset',headers=h,json={'confirmation':'RESET PYQ LIBRARY','cleanup_storage':False,'cancel_active':False})
    assert blocked.status_code==409
    done=client.post('/api/admin/library-reset',headers=h,json={'confirmation':'RESET PYQ LIBRARY','cleanup_storage':False,'cancel_active':True})
    assert done.status_code==200


def test_free_web_runner_processes_one_queued_file_per_request(app,client,monkeypatch):
    h=login(client,True);post_xlsx(client,'/api/admin/catalog/refresh',h,workbook())
    campaign=client.get('/api/admin/catalog/campaign',headers=h).json
    current=campaign['current']
    queued=client.post('/api/admin/catalog/campaign/process',headers=h,json={'stage':current['stage'],'term_id':current['term_id'],'limit':1})
    assert queued.status_code==202
    batch_id=queued.json['batch_id'];record=IngestionFile.query.filter_by(batch_id=batch_id).one()
    assert record.status=='FETCH_QUEUED'

    import backend.acquisition as acquisition
    import backend.ingestion as ingestion
    def fake_download(file_id):
        item=db.session.get(IngestionFile,file_id);item.status='QUEUED';item.path='fake.pdf';item.file_hash='a'*64;db.session.commit();return True
    def fake_process(file_id):
        item=db.session.get(IngestionFile,file_id);item.status='AVAILABLE';item.finished_at=123;db.session.commit();ingestion.update_batch(item.batch_id);return True
    monkeypatch.setattr(acquisition,'download_one',fake_download)
    monkeypatch.setattr(ingestion,'process_one',fake_process)

    step=client.post(f'/api/admin/catalog/batches/{batch_id}/run-next',headers=h,json={})
    assert step.status_code==200
    assert step.json['completed']==1
    assert step.json['queued']==0
    assert step.json['done'] is True
    assert step.json['percent']==100


def test_latest_batch_returns_resumable_queue(app,client):
    h=login(client,True);post_xlsx(client,'/api/admin/catalog/refresh',h,workbook())
    campaign=client.get('/api/admin/catalog/campaign',headers=h).json
    current=campaign['current']
    queued=client.post('/api/admin/catalog/campaign/process',headers=h,json={'stage':current['stage'],'term_id':current['term_id'],'limit':1})
    batch_id=queued.json['batch_id']
    latest=client.get('/api/admin/catalog/batches/latest',headers=h)
    assert latest.status_code==200
    assert latest.json['batch_id']==batch_id


def test_failed_term_does_not_block_next_and_remains_retryable(app,client):
    h=login(client,True);post_xlsx(client,'/api/admin/catalog/refresh',h,workbook())
    p=Paper.query.join(ExamType).join(Term).filter(ExamType.name=='Quiz 1',Term.name=='May 2026').one()
    b=IngestionBatch(user_id=User.query.filter_by(role='ADMIN').first().id);db.session.add(b);db.session.flush()
    db.session.add(IngestionFile(batch_id=b.id,paper_id=p.id,filename='failed.pdf',status='EXTRACTION_FAILED',source_url=p.source_url));db.session.commit()
    c=client.get('/api/admin/catalog/campaign',headers=h).json
    assert c['current']['term']=='Jan 2026'
    assert c['retry_target']['term']=='May 2026'
    assert c['failed']==1
    processed=client.post('/api/admin/catalog/campaign/process',headers=h,json={'limit':20})
    assert processed.status_code==202 and processed.json['queued']==1
    assert db.session.get(Paper,IngestionFile.query.filter_by(batch_id=processed.json['batch_id']).one().paper_id).term.name=='Jan 2026'
    retried=client.post('/api/admin/catalog/campaign/retry-failed',headers=h,json={'limit':20})
    assert retried.status_code==202 and retried.json['queued']==1
    assert IngestionFile.query.filter_by(batch_id=retried.json['batch_id']).one().paper_id==p.id


def test_admin_can_process_quiz2_before_quiz1(app,client):
    h=login(client,True);post_xlsx(client,'/api/admin/catalog/refresh',h,workbook())
    overview=client.get('/api/admin/catalog/campaign',headers=h).json
    target=next(g for g in overview['groups'] if g['stage']=='quiz2')
    result=client.post('/api/admin/catalog/campaign/process',headers=h,json={'stage':target['stage'],'term_id':target['term_id'],'limit':20})
    assert result.status_code==202 and result.json['queued']==1
    jobs=IngestionFile.query.all()
    assert len(jobs)==1
    assert db.session.get(Paper,jobs[0].paper_id).exam_type.name=='Quiz 2'
    assert result.json['current']['stage']=='quiz2'
    invalid=client.post('/api/admin/catalog/campaign/process',headers=h,json={'stage':'unknown','term_id':target['term_id'],'limit':20})
    assert invalid.status_code==400
    too_many=client.post('/api/admin/catalog/campaign/process',headers=h,json={'stage':'quiz1','term_id':target['term_id'],'limit':21})
    assert too_many.status_code==400


def test_recover_interrupted_file_preserves_completed_papers(app,client):
    import time
    h=login(client,True);post_xlsx(client,'/api/admin/catalog/refresh',h,workbook())
    target=client.get('/api/admin/catalog/campaign',headers=h).json['current']
    queued=client.post('/api/admin/catalog/campaign/process',headers=h,json={'stage':target['stage'],'term_id':target['term_id'],'limit':1})
    item=IngestionFile.query.filter_by(batch_id=queued.json['batch_id']).one()
    item.status='PROCESSING';item.started_at=time.time();db.session.commit()
    url=f'/api/admin/catalog/files/{item.id}/recover'
    assert client.post(url,headers=h,json={}).status_code==409
    item.started_at=time.time()-1900;db.session.commit()
    snapshot=client.get(f'/api/admin/catalog/batches/{item.batch_id}',headers=h).json
    assert snapshot['items'][0]['recoverable'] is True
    assert client.post(url,headers=h,json={}).status_code==200
    db.session.refresh(item);assert item.status=='PROCESSING_FAILED'
    assert client.post(url,headers=h,json={}).status_code==409
    assert Paper.query.count()>0
    retry=client.post(f'/api/admin/catalog/files/{item.id}/retry',headers=h,json={})
    assert retry.status_code==202 and retry.json['queued']==1
    assert client.post(f'/api/admin/catalog/files/{item.id}/retry',headers=h,json={}).json['queued']==0


def test_diploma_reset_preserves_other_levels_and_all_storage(app,client,monkeypatch):
    h=login(client,True)
    post_xlsx(client,'/api/admin/catalog/refresh',h,workbook())
    diploma=Course.query.first();diploma.level='Diploma'
    original=Paper.query.first()
    keep=Course(name='Degree course',code='KEEP',level='Degree');db.session.add(keep);db.session.flush()
    p=Paper(identity='keep-paper',name='Keep paper',course_id=keep.id,term_id=original.term_id,exam_type_id=original.exam_type_id)
    db.session.add(p);db.session.flush()
    q=Question(paper_id=p.id,number='1',kind='MCQ',text='Keep',status='AVAILABLE')
    removed=Question(paper_id=original.id,number='1',kind='MCQ',text='Remove',status='AVAILABLE')
    db.session.add_all([q,removed]);db.session.commit();keep_id=p.id;qid=q.id
    monkeypatch.setattr('backend.library_campaign.purge_project_prefix',lambda:(_ for _ in ()).throw(AssertionError('Storage must stay untouched')))
    preview=client.get('/api/admin/library-reset/diploma/preview',headers=h)
    assert preview.json['counts']['papers']==5 and preview.json['counts']['questions']==1
    assert client.post('/api/admin/library-reset/diploma',headers=h,json={'confirmation':'RESET PYQ LIBRARY'}).status_code==400
    done=client.post('/api/admin/library-reset/diploma',headers=h,json={'confirmation':'RESET DIPLOMA'})
    assert done.status_code==200 and done.json['storage_preserved']
    db.session.expire_all()
    assert Paper.query.count()==1 and db.session.get(Paper,keep_id)
    assert Question.query.count()==1 and db.session.get(Question,qid)
    assert Course.query.count()==2 and Term.query.count()==2
    assert post_xlsx(client,'/api/admin/catalog/refresh',h,workbook()).status_code==201
    assert Paper.query.count()==6


def test_diploma_reset_blocks_shared_bank_and_running_import(app,client):
    h=login(client,True);post_xlsx(client,'/api/admin/catalog/refresh',h,workbook())
    Course.query.first().level='Diploma';p=Paper.query.first()
    keep=Course(name='Foundation',code='FOUND',level='Foundation');db.session.add(keep);db.session.flush()
    alias=Paper(identity='shared-paper',name='Shared',course_id=keep.id,term_id=p.term_id,exam_type_id=p.exam_type_id,canonical_paper_id=p.id)
    db.session.add(alias);db.session.commit()
    endpoint='/api/admin/library-reset/diploma'
    assert client.post(endpoint,headers=h,json={'confirmation':'RESET DIPLOMA'}).status_code==409
    assert Paper.query.count()==6
    alias.canonical_paper_id=None
    batch=IngestionBatch(user_id=User.query.filter_by(role='ADMIN').first().id);db.session.add(batch);db.session.flush()
    db.session.add(IngestionFile(batch_id=batch.id,paper_id=p.id,filename='active.pdf',status='PROCESSING'))
    db.session.commit()
    assert client.post(endpoint,headers=h,json={'confirmation':'RESET DIPLOMA'}).status_code==409
    assert Paper.query.count()==6
