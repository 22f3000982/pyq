import io
import openpyxl
from backend.models import db,Course,Term,ExamType,Paper,Question,ImportRun,IngestionFile
from backend.catalog import paper_key
from conftest import login


def workbook_bytes(url='https://drive.google.com/file/d/source-one/view',paper_name='DL Quiz 1'):
    wb=openpyxl.Workbook()
    master=wb.active;master.title='Courses'
    master.append(['Course Name','Course Code','Course Level','Course Type'])
    master.append(['Deep Learning','DL','Diploma','Theory'])
    term=wb.create_sheet('Sep 2026')
    term.append(['Course Name','Quiz 1'])
    term.append(['Deep Learning',paper_name])
    term['B2'].hyperlink=url
    out=io.BytesIO();wb.save(out);wb.close();return out.getvalue()


def post_file(client,url,headers,data,extra=None):
    payload={'file':(io.BytesIO(data),'catalog.xlsx')}
    payload.update(extra or {})
    return client.post(url,headers=headers,data=payload,content_type='multipart/form-data')


def test_catalog_preview_is_non_mutating_and_apply_is_incremental(app,client):
    h=login(client,True);data=workbook_bytes()
    preview=post_file(client,'/api/admin/catalog/preview',h,data)
    assert preview.status_code==200
    assert preview.json['summary']['new']==1
    assert preview.json['summary']['available']==0
    assert Paper.query.count()==0 and ImportRun.query.count()==0

    applied=post_file(client,'/api/admin/catalog/apply',h,data,{
        'workbook_hash':preview.json['workbook_hash'],
        'process_new':'true','process_unprocessed':'true','changed_keys':'[]'
    })
    assert applied.status_code==202
    assert applied.json['new_papers']==1 and applied.json['queued']==1
    paper=Paper.query.one()
    assert paper.source_url=='https://drive.google.com/file/d/source-one/view'
    assert paper.status=='CATALOG_ONLY'  # queued backlog is not active processing
    assert IngestionFile.query.filter_by(paper_id=paper.id,status='FETCH_QUEUED').count()==1
    assert ImportRun.query.count()==1

    again=post_file(client,'/api/admin/catalog/preview',h,data)
    assert again.status_code==200 and again.json['already_applied'] is True
    assert again.json['summary']['new']==0
    assert again.json['summary']['unprocessed']==1
    assert Paper.query.count()==1


def test_changed_source_requires_explicit_selection(app,client):
    h=login(client,True)
    c=Course(name='Deep Learning',code='DL');t=Term(name='Sep 2026',year=2026,month=9);e=ExamType(name='Quiz 1')
    db.session.add_all([c,t,e]);db.session.flush()
    p=Paper(identity='existing-paper',course_id=c.id,term_id=t.id,exam_type_id=e.id,name='DL Quiz 1',
            source_url='https://drive.google.com/file/d/old-source/view',status='AVAILABLE')
    db.session.add(p);db.session.flush()
    db.session.add(Question(paper_id=p.id,number='1',kind='NAT',text='1+1?',answers=2,answer_status='ANSWER_AVAILABLE',marks=1,status='AVAILABLE'))
    db.session.commit()

    data=workbook_bytes('https://drive.google.com/file/d/new-source/view')
    preview=post_file(client,'/api/admin/catalog/preview',h,data)
    assert preview.json['summary']['changed']==1
    key=preview.json['items']['changed'][0]['key']

    keep=post_file(client,'/api/admin/catalog/apply',h,data,{
        'workbook_hash':preview.json['workbook_hash'],'process_new':'false','process_unprocessed':'false','changed_keys':'[]'
    })
    assert keep.status_code==202 and keep.json['updated_sources']==0 and keep.json['queued']==0
    assert db.session.get(Paper,p.id).source_url.endswith('/old-source/view')

    preview2=post_file(client,'/api/admin/catalog/preview',h,data)
    replace=post_file(client,'/api/admin/catalog/apply',h,data,{
        'workbook_hash':preview2.json['workbook_hash'],'process_new':'false','process_unprocessed':'false',
        'changed_keys':__import__('json').dumps([key])
    })
    assert replace.status_code==202 and replace.json['updated_sources']==1 and replace.json['queued']==1
    db.session.expire_all();updated=db.session.get(Paper,p.id)
    assert updated.status=='AVAILABLE'  # keep the working bank live until replacement actually processes
    assert updated.source_url.endswith('/new-source/view')
    assert Question.query.filter_by(paper_id=p.id,status='AVAILABLE').count()==1
    assert IngestionFile.query.filter_by(paper_id=p.id,status='FETCH_QUEUED').count()==1


def test_archived_paper_stays_ignored_and_missing_rows_never_delete(app,client):
    h=login(client,True)
    c=Course(name='Deep Learning',code='DL');t=Term(name='Sep 2026',year=2026,month=9);e=ExamType(name='Quiz 1')
    db.session.add_all([c,t,e]);db.session.flush()
    archived=Paper(identity='archived',course_id=c.id,term_id=t.id,exam_type_id=e.id,name='DL Quiz 1',
                   source_url='https://drive.google.com/file/d/source-one/view',status='ARCHIVED')
    extra=Paper(identity='keep-me',course_id=c.id,term_id=t.id,exam_type_id=e.id,name='Older paper',
                source_url='https://drive.google.com/file/d/older/view',status='CATALOG_ONLY')
    db.session.add_all([archived,extra]);db.session.commit()

    data=workbook_bytes()
    preview=post_file(client,'/api/admin/catalog/preview',h,data)
    assert preview.json['summary']['ignored']==1
    assert preview.json['summary']['absent']==1
    applied=post_file(client,'/api/admin/catalog/apply',h,data,{
        'workbook_hash':preview.json['workbook_hash'],'process_new':'true','process_unprocessed':'true','changed_keys':'[]'
    })
    assert applied.status_code==202
    db.session.expire_all()
    assert db.session.get(Paper,archived.id).status=='ARCHIVED'
    assert db.session.get(Paper,extra.id) is not None


def workbook_many(count=25):
    wb=openpyxl.Workbook()
    master=wb.active;master.title='Courses'
    master.append(['Course Name','Course Code','Course Level','Course Type'])
    master.append(['Deep Learning','DL','Diploma','Theory'])
    term=wb.create_sheet('Sep 2026')
    term.append(['Course Name','Quiz 1'])
    for n in range(1,count+1):
        row=n+1;term.cell(row,1,'Deep Learning');term.cell(row,2,f'DL QP{n}')
        term.cell(row,2).hyperlink=f'https://drive.google.com/file/d/source-{n}/view'
    out=io.BytesIO();wb.save(out);wb.close();return out.getvalue()


def test_preview_matches_existing_by_source_even_when_display_name_changed(app,client):
    h=login(client,True)
    c=Course(name='Deep Learning',code='DL');t=Term(name='Sep 2026',year=2026,month=9);e=ExamType(name='Quiz 1')
    db.session.add_all([c,t,e]);db.session.flush()
    p=Paper(identity='legacy-paper',course_id=c.id,term_id=t.id,exam_type_id=e.id,name='Old display title',
            source_url='https://drive.google.com/file/d/source-one/view',status='AVAILABLE')
    db.session.add(p);db.session.flush()
    db.session.add(Question(paper_id=p.id,number='1',kind='NAT',text='Value?',answers=1,answer_status='ANSWER_AVAILABLE',marks=1,status='AVAILABLE'))
    db.session.commit()
    preview=post_file(client,'/api/admin/catalog/preview',h,workbook_bytes(paper_name='DL Quiz 1 renamed'))
    assert preview.status_code==200
    assert preview.json['summary']['new']==0
    assert preview.json['summary']['absent']==0
    assert preview.json['summary']['available']==1


def test_sync_queues_only_requested_batch_and_next_preview_shrinks(app,client):
    h=login(client,True);data=workbook_many(25)
    preview=post_file(client,'/api/admin/catalog/preview',h,data)
    assert preview.status_code==200
    assert preview.json['summary']['new']==25
    applied=post_file(client,'/api/admin/catalog/apply',h,data,{
        'workbook_hash':preview.json['workbook_hash'],'process_new':'true','process_unprocessed':'true',
        'changed_keys':'[]','batch_limit':'20'
    })
    assert applied.status_code==202
    assert applied.json['queued']==20
    assert applied.json['selected_for_processing']==20
    assert applied.json['remaining_pending']==5
    assert Paper.query.count()==20
    assert IngestionFile.query.filter_by(status='FETCH_QUEUED').count()==20

    next_preview=post_file(client,'/api/admin/catalog/preview',h,data)
    assert next_preview.status_code==200
    assert next_preview.json['summary']['new']==5
    assert next_preview.json['summary']['queued']==20
    assert next_preview.json['summary']['pending']==5


def test_catalog_batch_progress_reports_percent_and_done(app,client):
    h=login(client,True);data=workbook_many(2)
    preview=post_file(client,'/api/admin/catalog/preview',h,data)
    applied=post_file(client,'/api/admin/catalog/apply',h,data,{
        'workbook_hash':preview.json['workbook_hash'],'process_new':'true','process_unprocessed':'true',
        'changed_keys':'[]','batch_limit':'2'
    })
    batch_id=applied.json['batch_id'];files=IngestionFile.query.filter_by(batch_id=batch_id).order_by(IngestionFile.id).all()
    files[0].status='AVAILABLE';files[1].status='EXTRACTION_FAILED';db.session.commit()
    progress=client.get(f'/api/admin/catalog/batches/{batch_id}',headers=h)
    assert progress.status_code==200
    assert progress.json['percent']==100
    assert progress.json['completed']==1
    assert progress.json['failed']==1
    assert progress.json['done'] is True
